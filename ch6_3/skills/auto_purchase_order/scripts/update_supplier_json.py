import os
import json
import argparse
from openpyxl import load_workbook

def update_supplier_json(xlsx_path, output_json_path):
    if not os.path.exists(xlsx_path):
        print(f"Error: Excel file '{xlsx_path}' not found.")
        return

    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    # Assuming first row is header
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    data = []

    for row in rows[1:]:
        if not any(row): continue # Skip empty rows
        entry = {header[i]: row[i] for i in range(len(header))}
        data.append(entry)

    with open(output_json_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f"Success: Updated {output_json_path} from {xlsx_path}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True, help="Path to input Excel file")
    parser.add_argument("--output", required=True, help="Path to output JSON file")
    args = parser.parse_args()
    
    update_supplier_json(args.xlsx, args.output)
