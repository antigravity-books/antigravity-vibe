import os
import json
import argparse
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage

def fill_purchase_order(purchase_json_path, supplier_name):
    # Paths (relative to scripts folder)
    base_dir = os.path.dirname(os.path.abspath(__file__))
    resource_dir = os.path.join(base_dir, "..", "resources")
    config_path = os.path.join(base_dir, "config.json")
    supplier_db_path = os.path.join(resource_dir, "구매협력업체정보.json")
    template_path = os.path.join(resource_dir, "발주서_엑셀양식.xlsx")

    # Load data
    with open(config_path, 'r', encoding='utf-8') as f:
        config = json.load(f)
    
    with open(supplier_db_path, 'r', encoding='utf-8') as f:
        suppliers = json.load(f)
    
    with open(purchase_json_path, 'r', encoding='utf-8') as f:
        purchase_data = json.load(f)

    # Find matching supplier
    supplier_info = next((s for s in suppliers if s['회사명'] == supplier_name), None)
    if not supplier_info:
        # Fallback: try partial match if not found exactly
        supplier_info = next((s for s in suppliers if supplier_name in s['회사명']), None)
    
    if not supplier_info:
        print(f"Error: Supplier '{supplier_name}' not found in DB.")
        return

    # Open Excel
    wb = load_workbook(template_path)
    ws = wb.active

    # 1. Fill Buyer Info (발주처)
    buyer = config['발주처']
    ws['D3'] = buyer['사업자등록번호']
    ws['D5'] = buyer['상호']
    ws['J5'] = buyer['담당자성명']
    ws['D6'] = buyer['사업장소재지']
    ws['D7'] = buyer['업태']
    ws['I7'] = buyer['종목']
    ws['D8'] = buyer['연락처']

    # 2. Fill Supplier Info (수주처)
    ws['R3'] = supplier_info['사업자등록번호']
    ws['R5'] = supplier_info['회사명']
    ws['X5'] = supplier_info['담당자']
    ws['R6'] = supplier_info['위치']
    ws['R7'] = supplier_info['업종'].split(' ')[0] if '업종' in supplier_info else "" # Simple mapping
    ws['W7'] = supplier_info['업종']
    ws['R8'] = supplier_info['연락처']

    # 3. Fill Items
    # Rows 11 to 26
    start_row = 11
    for i, item in enumerate(purchase_data):
        if i >= 16: break # Max 16 items
        row = start_row + i
        ws[f'A{row}'] = i + 1            # 순번
        ws[f'C{row}'] = f"{item['품목']} / {item['규격']}" # 품명/규격
        ws[f'H{row}'] = item['단위']      # 단위
        ws[f'K{row}'] = item['수량']      # 수량
        ws[f'N{row}'] = item['납기']      # 희망납품일자
        ws[f'R{row}'] = item['비고']      # 비고

    # 4. Set Date
    today = datetime.now()
    date_str = today.strftime("%Y년 %m월 %d일")
    ws['V28'] = f"일자 : {date_str}"

    # 5. Insert Signature
    signature_path = os.path.join(resource_dir, "signature.png")
    if os.path.exists(signature_path):
        img = OpenpyxlImage(signature_path)
        # Resize if necessary (signature space is around M5)
        img.width = 100
        img.height = 50
        ws.add_image(img, 'M5')
    
    # 6. Save File
    output_dir = os.path.join(base_dir, config['출력폴더'])
    os.makedirs(output_dir, exist_ok=True)
    
    filename = f"발주서_{today.strftime('%Y-%m-%d')}_{supplier_name}.xlsx".replace(" ", "_")
    output_path = os.path.join(output_dir, filename)
    wb.save(output_path)
    print(f"Success: Purchase order saved to {output_path}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--purchase", required=True, help="Path to purchase.json")
    parser.add_argument("--supplier", required=True, help="Name of the selected supplier")
    args = parser.parse_args()
    
    fill_purchase_order(args.purchase, args.supplier)
