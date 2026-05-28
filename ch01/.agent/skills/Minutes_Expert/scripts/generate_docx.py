#!/usr/bin/env python3
"""
Minutes_Expert AI Skill - Document Generator
템플릿 기반으로 회의록(.docx)을 생성하는 파이썬 스크립트입니다.

Requirements:
    pip install python-docx
"""

import sys
import os
try:
    from docx import Document
except ImportError:
    print("Error: 'python-docx' 라이브러리가 설치되어 있지 않습니다.")
    print("다음 명령어를 통해 설치해주세요: pip install python-docx")
    sys.exit(1)

def generate_minutes(template_path, output_path, meeting_data):
    """
    추출된 회의록 데이터를 바탕으로 템플릿의 플레이스홀더를 치환하여 새로운 docx 생성
    
    :param template_path: 템플릿 파일 경로
    :param output_path: 저장할 결과 파일 경로
    :param meeting_data: 치환할 데이터가 담긴 딕셔너리
    """
    if not os.path.exists(template_path):
        print(f"Error: 템플릿 파일을 찾을 수 없습니다: {template_path}")
        sys.exit(1)
        
    doc = Document(template_path)
    
    # 본문(단락) 치환
    for paragraph in doc.paragraphs:
        for key, value in meeting_data.items():
            placeholder = f"{{{{{key}}}}}"  # 예: {{DATE}}
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, str(value))
                
    # 표(Table) 내부 치환
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for key, value in meeting_data.items():
                        placeholder = f"{{{{{key}}}}}"
                        if placeholder in paragraph.text:
                            paragraph.text = paragraph.text.replace(placeholder, str(value))
                            
    doc.save(output_path)
    print(f"✅ 성공적으로 회의록을 생성했습니다: {output_path}")

if __name__ == "__main__":
    # AI 에이전트가 실행 시 활용할 수 있는 Mock Data 예시입니다.
    # 실제 환경에서는 arguments나 JSON 입력을 받아 meeting_data 딕셔너리를 구성해야 합니다.
    
    script_dir = os.path.dirname(os.path.abspath(__file__))
    template_file = os.path.join(script_dir, '..', 'resources', 'minutes_template.docx')
    output_file = os.path.join(script_dir, '..', 'generated_minutes.docx')
    
    # 템플릿에 맞춘 Mock Data (향후 AI가 이 부분을 동적으로 주입)
    mock_data = {
        "DATE": "2026-05-09",
        "LOCATION": "비대면 화상회의 (Zoom)",
        "ATTENDEES": "홍길동 팀장, 김철수 대리",
        "AGENDA": "신규 프로젝트 스킬 세팅",
        "DISCUSSION": "- 파이썬 스크립트 작성 여부 논의\n- 파일 복사 관련 자동화 논의",
        "DECISIONS": "[결정 01] scripts 내에 파이썬 코드 실제 구현\n[결정 02] 리소스 복사 스텝 명시",
        "ACTION_ITEMS": "완벽한 프롬프트 작성 (김철수, 05/10)"
    }
    
    generate_minutes(template_file, output_file, mock_data)
