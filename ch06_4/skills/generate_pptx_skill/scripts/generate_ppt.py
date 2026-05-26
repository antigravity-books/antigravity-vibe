import os
import sys
import copy
import json
import argparse
from pptx import Presentation

def apply_formatting(target_run, source_run):
    target_run.font.name = source_run.font.name
    if source_run.font.size is not None:
        target_run.font.size = source_run.font.size
    target_run.font.bold = source_run.font.bold
    target_run.font.italic = source_run.font.italic
    try:
        target_run.font.color.rgb = source_run.font.color.rgb
    except:
        pass

def main():
    parser = argparse.ArgumentParser(description="Generate PPTX using a pre-defined template and JSON replacement rules.")
    parser.add_argument("--input", required=True, help="Path to input JSON file mapping placeholders to text/images")
    parser.add_argument("--output", required=True, help="Path to save the generated PPTX")
    parser.add_argument("--template", default=r"..\resources\실습양식.pptx", help="Path to template PPTX")
    args = parser.parse_args()

    with open(args.input, 'r', encoding='utf-8') as f:
        data = json.load(f)

    text_replacements = data.get("text_replacements", {})
    image_replacements = data.get("image_replacements", {})

    os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)

    template_path = args.template
    if template_path.startswith(".."):
         template_path = os.path.join(os.path.dirname(__file__), template_path)
         template_path = os.path.abspath(template_path)

    prs = Presentation(template_path)

    for slide in prs.slides:
        for shape in list(slide.shapes):
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    p_text = paragraph.text
                    modified = False
                    
                    # Apply text replacement rules
                    for key, val_list in text_replacements.items():
                        if key in p_text and len(val_list) > 0:
                            if key in ["{{상세목차1}}", "{{상세목차2}}"]:
                                val = val_list.pop(0) if len(val_list) > 0 else val_list[-1]
                            else:
                                val = val_list[0]
                            p_text = p_text.replace(key, val)
                            modified = True
                    
                    # Apply image replacement rules
                    image_matched_key = None
                    for key in image_replacements.keys():
                        if key in paragraph.text:
                            image_matched_key = key
                            break
                            
                    if image_matched_key:
                        img_path = image_replacements[image_matched_key]
                        if os.path.exists(img_path):
                            x, y, cx, cy = shape.left, shape.top, shape.width, shape.height
                            slide.shapes.add_picture(img_path, x, y, cx, cy)
                            sp = shape._element
                            sp.getparent().remove(sp)
                            modified = False
                            break
                            
                    # Restore text formatting
                    elif modified:
                        if paragraph.runs:
                            first_run = paragraph.runs[0]
                            font_name = first_run.font.name
                            font_size = first_run.font.size
                            font_bold = first_run.font.bold
                            font_italic = first_run.font.italic
                            try:
                                font_color = first_run.font.color.rgb
                            except:
                                font_color = None
                                
                            alig = paragraph.alignment
                            paragraph.clear()
                            
                            new_run = paragraph.add_run()
                            new_run.text = p_text
                            if font_name is not None: new_run.font.name = font_name
                            if font_size is not None: new_run.font.size = font_size
                            if font_bold is not None: new_run.font.bold = font_bold
                            if font_italic is not None: new_run.font.italic = font_italic
                            if font_color is not None: new_run.font.color.rgb = font_color
                            paragraph.alignment = alig

    prs.save(args.output)

if __name__ == "__main__":
    main()
